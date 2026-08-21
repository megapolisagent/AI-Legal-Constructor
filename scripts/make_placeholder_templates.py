"""Тестовые заглушки шаблонов — НЕ юридический текст. Реальный текст готовят Мария/юрист
(handoff §9, «вне зоны Engineer»). Цель этого скрипта — только проверить, что движок
подстановки и сборки комплекта работает end-to-end. Запускать один раз при инициализации репо:
    python scripts/make_placeholder_templates.py
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "templates_docx"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def _doc_with_paragraphs(title: str, lines: list[str]) -> Document:
    doc = Document()
    doc.add_heading(title, level=1)
    for line in lines:
        doc.add_paragraph(line)
    return doc


def make_commission(deal_type: str) -> Document:
    return _doc_with_paragraphs(
        f"ДОГОВОР КОМИССИИ [ЗАГЛУШКА — {deal_type}]",
        [
            "[ТЕСТОВЫЙ ШАБЛОН. Юридический текст не согласован, использовать только для проверки движка.]",
            "Агентство: {{ agency_name }}, ИНН {{ agency_inn }}, ОГРН {{ agency_ogrn }}, в лице {{ agency_signatory }}.",
            "Участник: {{ participant.full_name }}, паспорт: {{ participant.passport_data }}, "
            "адрес регистрации: {{ participant.registration_address }}, контакт: {{ participant.contact }}.",
            "Сторона участника в сделке: {{ participant.side }}.",
            "Объект: {{ address }}, кадастровый номер {{ cadastral_number }}.",
            "Комиссионное вознаграждение: {{ commission_value }}.",
            "Тип сделки: {{ deal_type }}.",
        ],
    )


def make_main_contract(deal_type: str) -> Document:
    doc = _doc_with_paragraphs(
        f"ОСНОВНОЙ ДОГОВОР [ЗАГЛУШКА — {deal_type}]",
        [
            "[ТЕСТОВЫЙ ШАБЛОН. Юридический текст не согласован, использовать только для проверки движка.]",
            "Объект: {{ address }}, кадастровый номер {{ cadastral_number }}, площадь {{ area }}, "
            "комнат {{ rooms }}, основание права: {{ ownership_basis }}.",
            "Обременения: {{ encumbrances }}.",
            "Цена: {{ price }}. Порядок расчётов: {{ payment_schedule }}. Задаток/залог: {{ deposit }}.",
            "Срок: {{ term }}. Особые условия: {{ special_conditions }}.",
        ],
    )
    doc.add_heading("Сторона объекта:", level=2)
    doc.add_paragraph("{% for p in object_side %}{{ p.full_name }} (доля: {{ p.ownership_share }})\n{% endfor %}")
    doc.add_heading("Сторона контрагента:", level=2)
    doc.add_paragraph("{% for p in counterparty_side %}{{ p.full_name }}\n{% endfor %}")
    return doc


def make_act(deal_type: str) -> Document:
    doc = _doc_with_paragraphs(
        f"АКТ ПРИЁМА-ПЕРЕДАЧИ [ЗАГЛУШКА — {deal_type}]",
        [
            "[ТЕСТОВЫЙ ШАБЛОН. Юридический текст не согласован, использовать только для проверки движка.]",
            "Объект: {{ address }}, кадастровый номер {{ cadastral_number }}.",
            "Дата/срок передачи: {{ term }}.",
        ],
    )
    doc.add_heading("Участники сделки:", level=2)
    doc.add_paragraph("{% for p in participants %}{{ p.full_name }} — {{ p.side }}\n{% endfor %}")
    return doc


def make_consent(deal_type: str) -> Document:
    return _doc_with_paragraphs(
        f"СОГЛАСИЕ [ЗАГЛУШКА — {deal_type}]",
        [
            "[ТЕСТОВЫЙ ШАБЛОН. Юридический текст не согласован, использовать только для проверки движка.]",
            "[УТОЧНИТЬ у юриста: точный текст согласия зависит от основания — consent_reason={{ consent_reason }}]",
            "Участник, на которого оформлено согласие: {{ participant.full_name }}, "
            "паспорт: {{ participant.passport_data }}.",
            "Объект: {{ address }}.",
        ],
    )


BUILDERS = {
    "комиссия": make_commission,
    "основной_договор": make_main_contract,
    "акт": make_act,
    "согласие": make_consent,
}


def main() -> None:
    registry_entries = []
    for document_type, builder in BUILDERS.items():
        for deal_type in ("аренда", "продажа"):
            filename = f"{document_type}__{deal_type}__v1.docx"
            doc = builder(deal_type)
            doc.save(str(OUT_DIR / filename))
            registry_entries.append({
                "key": f"{document_type}__{deal_type}",
                "document_type": document_type,
                "deal_type": deal_type,
                "version": "v1",
                "status": "active",
                "file": filename,
            })
            print(f"создан {filename}")

    registry_path = OUT_DIR / "registry.json"
    registry_path.write_text(
        json.dumps({"entries": registry_entries}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"реестр записан в {registry_path}")


if __name__ == "__main__":
    main()
